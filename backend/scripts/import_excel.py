import os
import sys
import re
import json
from datetime import datetime, timezone
from collections import defaultdict

from openpyxl import load_workbook
from bson import json_util
from pymongo import ASCENDING

# Allow imports from the backend directory.
BACKEND_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

from db import get_collection


REQUIRED_SHEETS = {
    "Careers": ["Career ID", "Career name", "Career description"],
    "Skills": ["Skill ID", "Skill name", "Skill description"],
    "Chapters": ["Chapter ID", "Chapter name", "Chapter description", "Link"],
    "Career_Skill_Mapping": ["Career ID", "Skill ID", "Skill order"],
    "Skill_Chapter_Mapping": ["Skill ID", "Chapter ID", "Chapter order"],
}

COLLECTIONS = [
    "careers",
    "skills",
    "chapters",
    "career_skill_mapping",
    "skill_chapter_mapping",
]


def now_utc():
    return datetime.now(timezone.utc)


def cell_text(value):
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() in {"none", "nan"} else text


def business_id(value):
    # Preserve the workbook's ID spelling except for outer whitespace/case.
    return cell_text(value).upper()


def normalized_id(value):
    # Used only for matching/reporting, never to rewrite a stored business ID.
    return re.sub(r"\s+", "", business_id(value)).replace("O", "0")


def parse_order(value, sheet, row_number, issues, required=True):
    if value is None or cell_text(value) == "":
        if required:
            issues.append(f"{sheet} row {row_number}: order is empty.")
        return None
    try:
        number = float(value)
        if not number.is_integer():
            raise ValueError
        return int(number)
    except (ValueError, TypeError):
        issues.append(
            f"{sheet} row {row_number}: order {value!r} must be a whole number."
        )
        return None


def read_rows(workbook, sheet_name):
    ws = workbook[sheet_name]
    rows = []
    for row_number, row in enumerate(
        ws.iter_rows(min_row=2, values_only=True), start=2
    ):
        if not any(value is not None and cell_text(value) for value in row):
            continue
        rows.append((row_number, row))
    return rows


def validate_headers(workbook):
    problems = []
    for sheet_name, expected in REQUIRED_SHEETS.items():
        if sheet_name not in workbook.sheetnames:
            problems.append(f"Missing worksheet: {sheet_name}")
            continue
        ws = workbook[sheet_name]
        actual = [cell_text(v) for v in next(
            ws.iter_rows(min_row=1, max_row=1, values_only=True)
        )]
        actual = actual[:len(expected)]
        if actual != expected:
            problems.append(
                f"{sheet_name} headers should start with {expected!r}; "
                f"found {actual!r}."
            )
    return problems


def add_unique(records, identifier, label, sheet, row_number, issues):
    if not identifier:
        issues.append(f"{sheet} row {row_number}: {label} is empty.")
        return
    if identifier in records:
        issues.append(
            f"{sheet} row {row_number}: duplicate {label} {identifier!r} "
            f"(first seen at row {records[identifier]['row']})."
        )
        return
    records[identifier] = {"row": row_number}


def resolve_id(requested, exact_index, normalized_index):
    requested = business_id(requested)
    if not requested:
        return None, "missing"
    if requested in exact_index:
        return exact_index[requested], "exact"
    key = normalized_id(requested)
    if key not in normalized_index:
        return None, "missing"
    found = normalized_index[key]
    if found is None:
        return None, "ambiguous"
    return found, "normalized"


def prepare_workbook(workbook):
    """Read and validate the complete workbook without changing MongoDB."""
    issues = validate_headers(workbook)
    if issues:
        raise ValueError("\n".join(issues))

    report = {
        "created": {name: 0 for name in COLLECTIONS},
        "replaced": {name: 0 for name in COLLECTIONS},
        "issues": [],
        "resolvedIdAliases": [],
        "counts": {},
    }

    careers, skills, chapters = [], [], []
    career_by_id, skill_by_id, chapter_by_id = {}, {}, {}
    career_skill_rows, skill_chapter_rows = [], []

    # Careers
    for row_number, row in read_rows(workbook, "Careers"):
        cid = business_id(row[0] if len(row) > 0 else None)
        name = cell_text(row[1] if len(row) > 1 else None)
        description = cell_text(row[2] if len(row) > 2 else None)
        if not cid or not name:
            issues.append(
                f"Careers row {row_number}: Career ID and Career name are required."
            )
            continue
        if cid in career_by_id:
            issues.append(f"Careers row {row_number}: duplicate Career ID {cid!r}.")
            continue
        doc = {
            "career_id": cid,
            "name": name,
            "description": description,
            "active": True,
            "createdAt": now_utc(),
        }
        career_by_id[cid] = doc
        careers.append(doc)

    # Skills
    for row_number, row in read_rows(workbook, "Skills"):
        sid = business_id(row[0] if len(row) > 0 else None)
        name = cell_text(row[1] if len(row) > 1 else None)
        description = cell_text(row[2] if len(row) > 2 else None)
        if not sid or not name:
            issues.append(
                f"Skills row {row_number}: Skill ID and Skill name are required."
            )
            continue
        if sid in skill_by_id:
            issues.append(f"Skills row {row_number}: duplicate Skill ID {sid!r}.")
            continue
        doc = {
            "skill_id": sid,
            "name": name,
            "description": description,
            "active": True,
            "createdAt": now_utc(),
        }
        skill_by_id[sid] = doc
        skills.append(doc)

    # Chapters
    for row_number, row in read_rows(workbook, "Chapters"):
        chid = business_id(row[0] if len(row) > 0 else None)
        title = cell_text(row[1] if len(row) > 1 else None)
        description = cell_text(row[2] if len(row) > 2 else None)
        link = cell_text(row[3] if len(row) > 3 else None)
        if not chid or not title:
            issues.append(
                f"Chapters row {row_number}: Chapter ID and Chapter name are required."
            )
            continue
        if chid in chapter_by_id:
            issues.append(f"Chapters row {row_number}: duplicate Chapter ID {chid!r}.")
            continue
        doc = {
            "chapter_id": chid,
            "title": title,
            "description": description,
            "youtubeUrl": link,
            "order": len(chapters) + 1,
            "active": True,
            "createdAt": now_utc(),
        }
        chapter_by_id[chid] = doc
        chapters.append(doc)

    # Mapping source rows are held until all master IDs are known.
    for row_number, row in read_rows(workbook, "Career_Skill_Mapping"):
        cid = business_id(row[0] if len(row) > 0 else None)
        sid = business_id(row[1] if len(row) > 1 else None)
        order = parse_order(
            row[2] if len(row) > 2 else None,
            "Career_Skill_Mapping", row_number, issues
        )
        if not cid or not sid:
            issues.append(
                f"Career_Skill_Mapping row {row_number}: Career ID and Skill ID are required."
            )
            continue
        career_skill_rows.append((row_number, cid, sid, order))

    for row_number, row in read_rows(workbook, "Skill_Chapter_Mapping"):
        sid = business_id(row[0] if len(row) > 0 else None)
        chid = business_id(row[1] if len(row) > 1 else None)
        order = parse_order(
            row[2] if len(row) > 2 else None,
            "Skill_Chapter_Mapping", row_number, issues
        )
        if not sid or not chid:
            issues.append(
                f"Skill_Chapter_Mapping row {row_number}: Skill ID and Chapter ID are required."
            )
            continue
        skill_chapter_rows.append((row_number, sid, chid, order))

    # Build normalized lookup indexes. A normalized ID that corresponds to
    # more than one source record is deliberately marked ambiguous.
    def indexes(items, id_key):
        exact = {item[id_key]: item for item in items}
        groups = defaultdict(list)
        for key, item in exact.items():
            groups[normalized_id(key)].append(item)
        norm = {
            key: group[0] if len(group) == 1 else None
            for key, group in groups.items()
        }
        return exact, norm

    careers_exact, careers_norm = indexes(careers, "career_id")
    skills_exact, skills_norm = indexes(skills, "skill_id")
    chapters_exact, chapters_norm = indexes(chapters, "chapter_id")

    career_mapping_docs = []
    seen_career_pairs = set()
    for row_number, cid, sid, order in career_skill_rows:
        career, cres = resolve_id(cid, careers_exact, careers_norm)
        skill, sres = resolve_id(sid, skills_exact, skills_norm)
        if career is None:
            issues.append(
                f"Career_Skill_Mapping row {row_number}: Career ID {cid!r} "
                f"could not be resolved ({cres})."
            )
            continue
        if skill is None:
            issues.append(
                f"Career_Skill_Mapping row {row_number}: Skill ID {sid!r} "
                f"could not be resolved ({sres})."
            )
            continue
        pair = (career["career_id"], skill["skill_id"])
        if pair in seen_career_pairs:
            issues.append(
                f"Career_Skill_Mapping row {row_number}: duplicate mapping "
                f"{pair[0]} -> {pair[1]}."
            )
            continue
        seen_career_pairs.add(pair)
        # ObjectId references are filled after the master documents are inserted.
        career_mapping_docs.append({
            "_career_business_id": career["career_id"],
            "_skill_business_id": skill["skill_id"],
            "order": order,
            "source_row": row_number,
        })
        for requested, resolved, resolution, entity_type in (
            (cid, career["career_id"], cres, "career"),
            (sid, skill["skill_id"], sres, "skill"),
        ):
            if resolution == "normalized":
                report["resolvedIdAliases"].append({
                    "sheet": "Career_Skill_Mapping",
                    "row": row_number,
                    "entity": entity_type,
                    "excelId": requested,
                    "matchedId": resolved,
                })

    skill_mapping_docs = []
    seen_skill_pairs = set()
    for row_number, sid, chid, order in skill_chapter_rows:
        skill, sres = resolve_id(sid, skills_exact, skills_norm)
        chapter, chres = resolve_id(chid, chapters_exact, chapters_norm)
        if skill is None:
            issues.append(
                f"Skill_Chapter_Mapping row {row_number}: Skill ID {sid!r} "
                f"could not be resolved ({sres})."
            )
            continue
        if chapter is None:
            issues.append(
                f"Skill_Chapter_Mapping row {row_number}: Chapter ID {chid!r} "
                f"could not be resolved ({chres})."
            )
            continue
        pair = (skill["skill_id"], chapter["chapter_id"])
        if pair in seen_skill_pairs:
            issues.append(
                f"Skill_Chapter_Mapping row {row_number}: duplicate mapping "
                f"{pair[0]} -> {pair[1]}."
            )
            continue
        seen_skill_pairs.add(pair)
        skill_mapping_docs.append({
            "_skill_business_id": skill["skill_id"],
            "_chapter_business_id": chapter["chapter_id"],
            "order": order,
            "source_row": row_number,
        })
        if chres == "normalized":
            report["resolvedIdAliases"].append({
                "sheet": "Skill_Chapter_Mapping",
                "row": row_number,
                "entity": "chapter",
                "excelId": chid,
                "matchedId": chapter["chapter_id"],
                "note": "Matched using O/0 normalization; review the workbook ID.",
            })
        if sres == "normalized":
            report["resolvedIdAliases"].append({
                "sheet": "Skill_Chapter_Mapping",
                "row": row_number,
                "entity": "skill",
                "excelId": sid,
                "matchedId": skill["skill_id"],
                "note": "Matched using O/0 normalization; review the workbook ID.",
            })

    # Fail closed: never touch MongoDB if any validation issue exists.
    if issues:
        report["issues"] = issues
        return None, report

    report["counts"] = {
        "careers": len(careers),
        "skills": len(skills),
        "chapters": len(chapters),
        "career_skill_mapping": len(career_mapping_docs),
        "skill_chapter_mapping": len(skill_mapping_docs),
    }

    data = {
        "careers": careers,
        "skills": skills,
        "chapters": chapters,
        "career_skill_mapping": career_mapping_docs,
        "skill_chapter_mapping": skill_mapping_docs,
    }
    return data, report


def backup_collections(collection_names):
    """Save exact collection contents in MongoDB Extended JSON."""
    backup_dir = os.path.join(BACKEND_DIR, "backups")
    os.makedirs(backup_dir, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = os.path.join(backup_dir, f"master_data_backup_{stamp}.json")
    snapshot = {}
    for name in collection_names:
        snapshot[name] = list(get_collection(name).find({}))
    with open(backup_path, "w", encoding="utf-8") as file:
        file.write(json_util.dumps(snapshot, indent=2, ensure_ascii=False))
    return backup_path, snapshot


def replace_master_data(data, report):
    """Back up, then replace only the five master-data collections."""
    backup_path, snapshot = backup_collections(COLLECTIONS)

    careers = data["careers"]
    skills = data["skills"]
    chapters = data["chapters"]

    # Insert master documents first so mappings can reference their ObjectIds.
    try:
        for name in COLLECTIONS:
            collection = get_collection(name)
            collection.delete_many({})

        if careers:
            get_collection("careers").insert_many(careers)
        if skills:
            get_collection("skills").insert_many(skills)
        if chapters:
            get_collection("chapters").insert_many(chapters)

        career_index = {
            doc["career_id"]: doc["_id"]
            for doc in get_collection("careers").find(
                {}, {"_id": 1, "career_id": 1}
            )
        }
        skill_index = {
            doc["skill_id"]: doc["_id"]
            for doc in get_collection("skills").find(
                {}, {"_id": 1, "skill_id": 1}
            )
        }
        chapter_index = {
            doc["chapter_id"]: doc["_id"]
            for doc in get_collection("chapters").find(
                {}, {"_id": 1, "chapter_id": 1}
            )
        }

        career_mappings = []
        for item in data["career_skill_mapping"]:
            career_mappings.append({
                "careerId": career_index[item["_career_business_id"]],
                "skillId": skill_index[item["_skill_business_id"]],
                "order": item["order"],
                "createdAt": now_utc(),
            })

        skill_mappings = []
        for item in data["skill_chapter_mapping"]:
            skill_mappings.append({
                "skillId": skill_index[item["_skill_business_id"]],
                "chapterId": chapter_index[item["_chapter_business_id"]],
                "order": item["order"],
                "createdAt": now_utc(),
            })

        if career_mappings:
            get_collection("career_skill_mapping").insert_many(career_mappings)
        if skill_mappings:
            get_collection("skill_chapter_mapping").insert_many(skill_mappings)

        # Basic post-import count check.
        for name in COLLECTIONS:
            actual = get_collection(name).count_documents({})
            expected = report["counts"][name]
            if actual != expected:
                raise RuntimeError(
                    f"Post-import count mismatch for {name}: "
                    f"expected {expected}, found {actual}."
                )

        report["backupFile"] = backup_path
        report["replaced"] = {
            name: report["counts"][name] for name in COLLECTIONS
        }
        return report

    except Exception:
        # Best-effort rollback from the in-memory snapshot if replacement fails.
        for name in COLLECTIONS:
            collection = get_collection(name)
            collection.delete_many({})
            old_docs = snapshot.get(name, [])
            if old_docs:
                collection.insert_many(old_docs)
        raise


def print_report(report):
    print("\n========== CAREERCRAFT MASTER DATA IMPORT ==========")
    if report.get("issues"):
        print("\nVALIDATION FAILED — DATABASE WAS NOT CHANGED.")
        for issue in report["issues"]:
            print(f" - {issue}")
        print(f"\nTotal issues: {len(report['issues'])}")
        return

    print("\nImported document counts:")
    for name, count in report.get("replaced", {}).items():
        print(f"  {name}: {count}")

    print(f"\nBackup file: {report.get('backupFile', 'not available')}")
    aliases = report.get("resolvedIdAliases", [])
    print(f"Resolved O/0 or normalized aliases: {len(aliases)}")
    for item in aliases:
        print(
            f"  [{item['sheet']} row {item['row']}] "
            f"{item['excelId']} -> {item['matchedId']}"
        )
    print("\nImport completed successfully.")


def main():
    if len(sys.argv) < 2:
        print(
            "Usage: python scripts/import_excel.py "
            "data/UPDATEDOCRCFT(1).xlsx"
        )
        sys.exit(2)

    file_path = os.path.abspath(sys.argv[1])
    if not os.path.isfile(file_path):
        print(f"File not found: {file_path}")
        sys.exit(2)

    workbook = None
    try:
        workbook = load_workbook(
            file_path,
            data_only=True,
            read_only=True,
        )
        data, report = prepare_workbook(workbook)

        if data is None:
            print_report(report)
            sys.exit(1)

        print("Workbook validation passed.")
        print("Records to replace:")
        for name, count in report["counts"].items():
            print(f"  {name}: {count}")

        # Explicit confirmation protects against accidental replacement.
        confirmation = input(
            "\nThis will replace ONLY the five master-data collections. "
            "A backup will be created first. Type REPLACE to continue: "
        ).strip()
        if confirmation != "REPLACE":
            print("Cancelled. MongoDB was not changed.")
            return

        final_report = replace_master_data(data, report)
        print_report(final_report)

    except Exception as exc:
        print(f"\nIMPORT FAILED: {exc}")
        print("If replacement had started, the script attempted a rollback.")
        sys.exit(1)
    finally:
        if workbook is not None:
            workbook.close()


if __name__ == "__main__":
    main()
